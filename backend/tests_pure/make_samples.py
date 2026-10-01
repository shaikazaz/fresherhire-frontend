"""Generates sample_data/sample_resume.pdf, sample_resume.docx and weak_resume.pdf (needs reportlab + python-docx)."""
import os

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = os.path.join(os.path.dirname(__file__), "..", "sample_data")

GOOD = [
    ("Ananya Reddy", None),
    ("Hyderabad, Telangana | ananya.reddy@example.com | +91 98765 43210", None),
    ("linkedin.com/in/ananyareddy | github.com/ananyareddy", None),
    ("Education", "H"),
    ("B.Tech in Computer Science and Engineering, JNTU Hyderabad, 2021 - 2025, CGPA 8.4", None),
    ("Skills", "H"),
    ("Python, Django, JavaScript, React, SQL, MySQL, HTML, CSS, Git, REST API, Bootstrap, Docker", None),
    ("Projects", "H"),
    ("Job Portal Platform: Built a job marketplace using Django and MySQL with role based login and REST API.", None),
    ("Designed a resume scanner in Python that analyzed 200+ resumes and improved keyword coverage by 35%.", None),
    ("Developed a React dashboard with charts and filters used by 50 students during testing.", None),
    ("Expense Tracker: Implemented an expense tracking app with JavaScript, HTML and CSS; reduced entry time by 40%.", None),
    ("Internships", "H"),
    ("Python Developer Intern, TechNova Solutions, Jun 2024 - Aug 2024", None),
    ("Automated weekly reports with Python and SQL, saving 5 hours per week for a team of 8.", None),
    ("Collaborated with 4 engineers to test and deploy Django features using Git and Docker.", None),
    ("Certifications", "H"),
    ("Python for Everybody, Coursera; SQL Fundamentals, Udemy", None),
]

WEAK = [
    ("Ravi Kumar", None),
    ("I am a hardworking student looking for a job in a good company where I can learn.", None),
    ("Hobbies", "H"),
    ("Cricket, movies, travelling", None),
]


def make_pdf(path, lines):
    c = canvas.Canvas(path, pagesize=A4)
    y = 800
    for text, kind in lines:
        c.setFont("Helvetica-Bold" if kind == "H" else "Helvetica", 12 if kind == "H" else 10)
        if kind == "H":
            y -= 8
        c.drawString(50, y, text[:105])
        y -= 16
    c.save()


def make_docx(path, lines):
    d = Document()
    for text, kind in lines:
        if kind == "H":
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(text)
    d.save(path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    make_pdf(os.path.join(OUT, "sample_resume.pdf"), GOOD)
    make_docx(os.path.join(OUT, "sample_resume.docx"), GOOD)
    make_pdf(os.path.join(OUT, "weak_resume.pdf"), WEAK)
    print("samples written to", os.path.abspath(OUT))

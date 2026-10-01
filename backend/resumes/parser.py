"""Extract text and profile data from PDF / DOCX resumes.
No Django imports.
"""

import math
import re

from .ats import EMAIL_RE, find_phone, split_sections


# ============================================================
# REGEX PATTERNS
# ============================================================

EDU_RE = re.compile(
    r"\b("
    r"B\.?\s?Tech|"
    r"B\.?\s?E\b|"
    r"B\.?\s?Sc|"
    r"BCA|"
    r"BBA|"
    r"B\.?\s?Com|"
    r"M\.?\s?Tech|"
    r"M\.?\s?Sc|"
    r"MCA|"
    r"MBA|"
    r"BE\b|"
    r"ME\b|"
    r"Bachelor[^\n,]{0,80}|"
    r"Master[^\n,]{0,80}|"
    r"Diploma[^\n,]{0,80}"
    r")",
    re.I,
)


YEAR_RE = re.compile(
    r"\b(19\d{2}|20\d{2})\b"
)


# ============================================================
# PDF EXTRACTION
# ============================================================

def _pdf(file_obj):
    import pdfplumber

    texts = []
    tables = 0
    images = 0

    with pdfplumber.open(file_obj) as pdf:

        pages = len(pdf.pages)

        for page in pdf.pages:

            texts.append(
                page.extract_text() or ""
            )

            images += len(page.images)

            try:

                for table in page.find_tables():

                    rows = table.extract()

                    if (
                        len(rows) >= 2
                        and max(
                            len(r)
                            for r in rows
                        ) >= 2
                    ):
                        tables += 1

            except Exception:
                pass

    return {
        "text": "\n".join(texts),
        "pages": pages,
        "tables": tables,
        "images": images,
    }


# ============================================================
# DOCX EXTRACTION
# ============================================================

def _docx(file_obj):
    from docx import Document

    doc = Document(file_obj)

    parts = []

    # Normal paragraphs
    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in doc.tables:

        for row in table.rows:

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    parts.append(text)

    text = "\n".join(parts)

    words = len(text.split())

    return {
        "text": text,
        "pages": max(
            1,
            math.ceil(words / 450)
        ),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
    }


# ============================================================
# MAIN TEXT EXTRACTION
# ============================================================

def extract_text(file_obj, ext):

    ext = ext.lower().lstrip(".")

    if ext == "pdf":
        return _pdf(file_obj)

    if ext == "docx":
        return _docx(file_obj)

    raise ValueError(
        "Unsupported file type: " + ext
    )


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_line(line):
    """Clean a single resume line."""

    line = line.strip()

    line = re.sub(
        r"\s+",
        " ",
        line,
    )

    return line


def clean_phone(phone):
    """Clean extracted phone number."""

    if not phone:
        return ""

    phone = str(phone).strip()

    # Remove common labels
    phone = re.sub(
        r"^(phone|mobile|mob|contact|tel)\s*[:\-]?\s*",
        "",
        phone,
        flags=re.I,
    )

    phone = re.sub(
        r"\s+",
        " ",
        phone,
    )

    return phone[:30]


# ============================================================
# NAME EXTRACTION
# ============================================================

def extract_full_name(text):
    """
    Best-effort extraction of candidate full name.

    Usually the candidate name is near the top of the resume.
    """

    lines = []

    for line in text.splitlines():

        line = clean_line(line)

        if not line:
            continue

        lines.append(line)

    if not lines:
        return ""

    ignored_words = {
        "resume",
        "curriculum vitae",
        "cv",
        "profile",
        "summary",
        "objective",
        "contact",
        "phone",
        "mobile",
        "email",
        "linkedin",
        "github",
        "portfolio",
        "education",
        "experience",
        "skills",
        "projects",
        "certifications",
        "achievements",
        "internship",
        "internships",
        "work experience",
        "professional summary",
    }

    # Only inspect the top portion.
    for line in lines[:20]:

        lower = line.lower().strip()

        if lower in ignored_words:
            continue

        # Email
        if EMAIL_RE.search(line):
            continue

        # URLs
        if re.search(
            r"https?://|www\.|linkedin\.com|github\.com",
            lower,
        ):
            continue

        # Phone
        if find_phone(line):
            continue

        # Lines containing numbers are usually not names
        if re.search(r"\d", line):
            continue

        # Remove bullets
        line = re.sub(
            r"^[•●▪◦\-*]+\s*",
            "",
            line,
        ).strip()

        words = line.split()

        # Candidate name usually has 2-5 words
        if not 2 <= len(words) <= 5:
            continue

        valid = True

        for word in words:

            # Allow names such as:
            # Shaik Azaz
            # A. Kumar
            # John O'Connor
            # Mary-Jane Smith
            if not re.fullmatch(
                r"[A-Za-z.'-]+",
                word,
            ):
                valid = False
                break

        if valid:

            # Reject obvious headings
            if all(
                word.lower()
                not in ignored_words
                for word in words
            ):

                return line[:150]

    return ""


# ============================================================
# PHONE EXTRACTION
# ============================================================

def extract_phone(text):
    """
    Extract phone number from the complete resume.
    Uses the project's existing find_phone() first,
    then falls back to common Indian/international formats.
    """

    # Existing project phone extractor
    try:

        phone = find_phone(text)

        if phone:
            return clean_phone(phone)

    except Exception:
        pass

    # Common fallback patterns
    patterns = [

        # +91 9876543210
        r"\+91[\s\-]?[6-9]\d{9}",

        # 91 9876543210
        r"\b91[\s\-]?[6-9]\d{9}\b",

        # 98765 43210
        r"\b[6-9]\d{4}[\s\-]\d{5}\b",

        # 9876543210
        r"\b[6-9]\d{9}\b",

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
        )

        if match:

            return clean_phone(
                match.group(0)
            )

    return ""


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

def extract_education(text, sections):
    """
    Extract the most relevant education line.
    """

    education_section = sections.get(
        "education",
        "",
    )

    # Prefer Education section
    candidates = []

    if education_section:
        candidates.extend(
            education_section.splitlines()
        )

    # If Education section is not detected,
    # search the entire resume.
    candidates.extend(
        text.splitlines()
    )

    for line in candidates:

        line = clean_line(line)

        if not line:
            continue

        if EDU_RE.search(line):

            # Avoid extremely long lines
            return line[:200]

    return ""


# ============================================================
# GRADUATION YEAR EXTRACTION
# ============================================================

def extract_grad_year(
    text,
    sections,
):
    """
    Extract graduation year.

    Prefer the Education section.
    Accepts years between 1990 and 2040.
    """

    education_text = sections.get(
        "education",
        "",
    )

    # Search education section first
    years = [
        int(year)
        for year in YEAR_RE.findall(
            education_text
        )
    ]

    years = [
        year
        for year in years
        if 1990 <= year <= 2040
    ]

    if years:

        # Usually the latest year is graduation year
        return max(years)

    # Fallback to entire resume
    years = [
        int(year)
        for year in YEAR_RE.findall(
            text
        )
    ]

    years = [
        year
        for year in years
        if 1990 <= year <= 2040
    ]

    if years:

        return max(years)

    return None


# ============================================================
# LOCATION EXTRACTION
# ============================================================

def extract_location(text):
    """
    Best-effort city/location extraction.

    This intentionally avoids guessing aggressively.
    """

    location_patterns = [
        r"(?:location|city|address)\s*[:\-]\s*([A-Za-z .,'-]{2,100})",
        r"(?:based in|located in)\s+([A-Za-z .,'-]{2,100})",
    ]

    for pattern in location_patterns:

        match = re.search(
            pattern,
            text,
            flags=re.I,
        )

        if match:

            value = clean_line(
                match.group(1)
            )

            if value:

                return value[:120]

    return ""


# ============================================================
# EXPERIENCE LEVEL
# ============================================================

def extract_experience_level(text):
    """
    Best-effort experience-level detection.
    """

    lower = text.lower()

    # Fresher
    fresher_words = [
        "fresher",
        "recent graduate",
        "recently graduated",
        "entry level",
        "entry-level",
        "graduate",
    ]

    for word in fresher_words:

        if word in lower:
            return "fresher"

    # Junior
    junior_patterns = [
        r"\b0\s*[-–]?\s*2\s+years?\b",
        r"\b1\s+year\b",
        r"\b2\s+years?\b",
    ]

    for pattern in junior_patterns:

        if re.search(
            pattern,
            lower,
        ):
            return "junior"

    # Mid
    mid_patterns = [
        r"\b2\s*[-–]?\s*5\s+years?\b",
        r"\b3\s+years?\b",
        r"\b4\s+years?\b",
        r"\b5\s+years?\b",
    ]

    for pattern in mid_patterns:

        if re.search(
            pattern,
            lower,
        ):
            return "mid"

    return ""


# ============================================================
# MAIN PROFILE EXTRACTION
# ============================================================

def extract_profile(text):
    """
    Extract profile information from resume text.

    Returns:
        full_name
        phone
        email
        education
        grad_year
        location
        experience_level
        projects_text
    """

    # Safety
    if not text:
        return {
            "full_name": "",
            "phone": "",
            "email": "",
            "education": "",
            "grad_year": None,
            "location": "",
            "experience_level": "",
            "projects_text": "",
        }

    # Split resume into sections
    sections = split_sections(text)

    # --------------------------------------------------------
    # Name
    # --------------------------------------------------------

    full_name = extract_full_name(
        text
    )

    # --------------------------------------------------------
    # Phone
    # --------------------------------------------------------

    phone = extract_phone(
        text
    )

    # --------------------------------------------------------
    # Email
    # --------------------------------------------------------

    email_match = EMAIL_RE.search(
        text
    )

    email = (
        email_match.group(0).strip()
        if email_match
        else ""
    )

    # --------------------------------------------------------
    # Education
    # --------------------------------------------------------

    education = extract_education(
        text,
        sections,
    )

    # --------------------------------------------------------
    # Graduation year
    # --------------------------------------------------------

    grad_year = extract_grad_year(
        text,
        sections,
    )

    # --------------------------------------------------------
    # Location
    # --------------------------------------------------------

    location = extract_location(
        text
    )

    # --------------------------------------------------------
    # Experience level
    # --------------------------------------------------------

    experience_level = (
        extract_experience_level(text)
    )

    # --------------------------------------------------------
    # Projects
    # --------------------------------------------------------

    projects_text = sections.get(
        "projects",
        "",
    )

    # --------------------------------------------------------
    # Return everything
    # --------------------------------------------------------

    return {
        "full_name": full_name,
        "phone": phone,
        "email": email,
        "education": education,
        "grad_year": grad_year,
        "location": location,
        "experience_level": experience_level,
        "projects_text": projects_text,
    }